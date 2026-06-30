from docx import Document

from src.models.profile import Profile


class CVBuilder:

    def __init__(self, profile: Profile):
        self.profile = profile

    def build(self):

        document = Document()

        document.add_heading(self.profile.full_name, level=1)

        document.add_paragraph(self.profile.professional_title)

        document.add_paragraph(self.profile.email)

        document.add_paragraph(self.profile.phone)

        document.add_paragraph(self.profile.location)

        document.save("Kenleigh_Peters_CV.docx")
        